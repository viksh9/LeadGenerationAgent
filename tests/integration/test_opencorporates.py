"""OpenCorporates persistence + API + full-intelligence export (Prompt 47).

Isolated DB. OpenCorporates is driven by a mocked client (no network). Verifies legal
fields, registered address kept DISTINCT from operating address, officers separate
from POCs, conflict retention, Data Trust, NO_MATCH (no fabrication), the API, the
full-intelligence export, and that the fixed 16-column export is unchanged.
"""

from __future__ import annotations

from io import BytesIO

import httpx
import openpyxl
import pytest

import api.main
from api.dependencies import get_session
from config import get_settings
from database.models import (
    Company,
    CompanyFieldEvidence,
    CompanyLocation,
    CompanyOfficer,
    DataProvenance,
    Lead,
    LeadPriority,
    LeadStatus,
    SignalType,
)
from enrichment.public_intelligence_service import discover_company_intelligence
from integrations.public_intelligence.opencorporates.client import OpenCorporatesClient
from integrations.public_intelligence.opencorporates.provider import OpenCorporatesProvider

_SEARCH = {"results": {"total_count": 1, "companies": [{"company": {
    "name": "Infosys Limited", "company_number": "L85110KA1981PLC013115", "jurisdiction_code": "in_ka",
    "current_status": "Active", "registry_url": "https://mca.gov.in/x",
    "opencorporates_url": "https://opencorporates.com/companies/in_ka/L85110KA1981PLC013115"}}]}}
_DETAIL = {"results": {"company": {
    "name": "Infosys Limited", "company_number": "L85110KA1981PLC013115", "jurisdiction_code": "in_ka",
    "current_status": "Active", "incorporation_date": "1981-07-02",
    "registry_url": "https://mca.gov.in/x",
    "opencorporates_url": "https://opencorporates.com/companies/in_ka/L85110KA1981PLC013115",
    "registered_address": {"street_address": "Plot 44, Electronics City", "locality": "Bengaluru",
                           "region": "Karnataka", "postal_code": "560100", "country": "India"},
    "officers": [{"officer": {"name": "Salil Parekh", "position": "CEO & MD"}}],
    "source": {"publisher": "MCA"}}}}


@pytest.fixture(autouse=True)
def _token():
    s = get_settings(); prev = s.opencorporates_api_token
    s.opencorporates_api_token = "tok"
    yield
    s.opencorporates_api_token = prev


def _oc_provider(handler=None):
    handler = handler or (lambda r: httpx.Response(200, json=(_SEARCH if "/search" in r.url.path else _DETAIL)))
    client = OpenCorporatesClient(http=httpx.Client(transport=httpx.MockTransport(handler)), sleep=lambda s: None)
    return OpenCorporatesProvider(client=client)


def _company(session, *, name="Infosys", domain="infosys.com", india=True, website="https://infosys.com"):
    c = Company(canonical_name=name, normalized_name=name.lower(), primary_domain=domain, website=website,
                full_address="Bengaluru, Karnataka, India", headquarters_city="Bengaluru",
                india_presence=india, data_provenance=DataProvenance.REAL)
    session.add(c)
    session.commit()
    return c


def test_persists_legal_fields_and_registered_address_distinct(seed_session):
    company = _company(seed_session)
    summ = discover_company_intelligence(seed_session, company, providers=[_oc_provider()])
    seed_session.refresh(company)
    assert company.legal_name == "Infosys Limited" and company.company_number == "L85110KA1981PLC013115"
    assert company.jurisdiction_code == "in_ka" and company.company_status == "ACTIVE"
    assert company.india_entity_type == "INDIA_ENTITY"
    # Operating address (website) preserved; registered address stored separately (§9).
    assert company.full_address == "Bengaluru, Karnataka, India"
    assert "Plot 44" in (company.registered_address or "")
    locs = seed_session.query(CompanyLocation).filter_by(company_id=company.id).all()
    types = {l.location_type for l in locs}
    assert "REGISTERED_OFFICE" in types
    assert summ.data_trust_score >= 30   # registry + OC match contribute


def test_officers_separate_from_pocs(seed_session):
    company = _company(seed_session)
    discover_company_intelligence(seed_session, company, providers=[_oc_provider()])
    officers = seed_session.query(CompanyOfficer).filter_by(company_id=company.id).all()
    assert officers and officers[0].name == "Salil Parekh" and officers[0].role_kind == "LEGAL_OFFICER"


def test_no_match_no_fabrication(seed_session):
    company = _company(seed_session, name="Nonexistent Co", domain=None, website=None)
    provider = _oc_provider(lambda r: httpx.Response(200, json={"results": {"total_count": 0, "companies": []}}))
    discover_company_intelligence(seed_session, company, providers=[provider])
    seed_session.refresh(company)
    assert company.legal_name is None and company.company_number is None   # nothing invented


def test_conflicting_address_retained(seed_session):
    company = _company(seed_session)
    discover_company_intelligence(seed_session, company, providers=[_oc_provider()])
    seed_session.add(CompanyFieldEvidence(
        company_id=company.id, field="registered_address", value="Different, Mumbai",
        source="Official Company Website", source_type="company_website",
        source_url="https://infosys.com/contact", source_priority=1, trust_score=95,
        data_provenance=DataProvenance.REAL))
    seed_session.commit()
    rows = seed_session.query(CompanyFieldEvidence).filter_by(
        company_id=company.id, field="registered_address").all()
    assert len(rows) == 2   # both sources retained (§22)


# --------------------------------------------------------------------------- #
# API + export
# --------------------------------------------------------------------------- #
def _seed_via_client(client):
    session = next(api.main.app.dependency_overrides[get_session]())
    c = Company(canonical_name="Infosys", normalized_name="infosys", primary_domain="infosys.com",
                website="https://infosys.com", legal_name="Infosys Limited",
                company_number="L85110KA1981PLC013115", jurisdiction_code="in_ka",
                company_status="ACTIVE", registered_address="Plot 44, Electronics City, Bengaluru",
                full_address="Bengaluru, Karnataka, India", headquarters_city="Bengaluru",
                india_presence=True, india_entity_type="INDIA_ENTITY", data_trust_score=100,
                registry_url="https://mca.gov.in/x",
                opencorporates_url="https://opencorporates.com/companies/in_ka/L85110KA1981PLC013115",
                data_provenance=DataProvenance.REAL)
    session.add(c)
    session.flush()
    session.add(CompanyOfficer(company_id=c.id, name="Salil Parekh", normalized_name="salil parekh",
                               position="CEO & MD", role_kind="LEGAL_OFFICER", source="OpenCorporates",
                               data_provenance=DataProvenance.REAL))
    session.add(Lead(company_name="Infosys", normalized_company_name="infosys", lead_score=85,
                     lead_priority=LeadPriority.HOT, status=LeadStatus.NEW, signal_type=SignalType.HIRING,
                     estimated_hiring=40, it_job_count=40, technologies=["Java"],
                     data_provenance=DataProvenance.REAL))
    cid = c.id
    session.commit()
    session.close()
    return cid


def test_api_profile_has_legal_fields(client):
    cid = _seed_via_client(client)
    body = client.get(f"/companies/{cid}/public-intelligence").json()
    assert body["legal_name"] == "Infosys Limited" and body["company_number"] == "L85110KA1981PLC013115"
    assert body["company_status"] == "ACTIVE" and body["india_entity_type"] == "INDIA_ENTITY"
    assert body["registered_address"] == "Plot 44, Electronics City, Bengaluru"
    assert len(body["officers"]) == 1 and body["officers"][0]["role_kind"] == "LEGAL_OFFICER"


def test_api_opencorporates_status_no_secret(client):
    body = client.get("/integrations/opencorporates/status").json()
    assert body["status"] in ("CONFIGURED", "NOT_CONFIGURED", "DISABLED")
    assert "token" not in body and "api_token" not in body
    assert body["api_version"] == "v0.4"


def test_16_column_export_unchanged(client):
    _seed_via_client(client)
    r = client.get("/export/excel?scope=all")
    wb = openpyxl.load_workbook(BytesIO(r.content), read_only=True)
    assert wb.sheetnames == ["Lead Data"]
    headers = [c.value for c in next(wb["Lead Data"].iter_rows(min_row=1, max_row=1))]
    assert len(headers) == 16   # fixed business contract preserved


def test_full_intelligence_export(client):
    _seed_via_client(client)
    r = client.get("/export/full-intelligence")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.content), read_only=True)
    ws = wb["Full Intelligence"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    row = dict(zip(headers, list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]))
    assert row["Legal Company Name"] == "Infosys Limited"
    assert row["Company Number"] == "L85110KA1981PLC013115"
    assert row["Operating Address"] == "Bengaluru, Karnataka, India"
    assert row["Registered Address"] == "Plot 44, Electronics City, Bengaluru"   # distinct
    assert row["India Entity Type"] == "INDIA_ENTITY" and row["Data Trust"] == 100
