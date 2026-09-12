"""Full-intelligence export — enrichment columns (Prompt 49, §39/§40).

Verifies the extended full-intelligence workbook carries the new enrichment fields
(Role Match, Email/Phone Verification) sourced from real persisted POC values, and that
the FIXED 16-column export remains untouched.
"""

from __future__ import annotations

from io import BytesIO

import openpyxl

from database.models import (
    Company,
    ContactType,
    DataProvenance,
    DecisionMaker,
    EmailStatus,
    Lead,
    LeadPriority,
    LeadStatus,
    SignalType,
    VerificationStatus,
)
from export.excel import HEADERS as FIXED_HEADERS
from export.full_intelligence import build_full_intelligence_workbook


def _seed(session):
    session.add(Company(canonical_name="Acme Corp", normalized_name="acme corp",
                        primary_domain="acme.com", data_provenance=DataProvenance.REAL))
    lead = Lead(company_name="Acme Corp", normalized_company_name="acme corp", lead_score=80,
                lead_priority=LeadPriority.HOT, status=LeadStatus.NEW, signal_type=SignalType.HIRING,
                estimated_hiring=30, data_provenance=DataProvenance.REAL)
    session.add(lead)
    session.flush()
    company = session.query(Company).first()
    session.add(DecisionMaker(
        company_id=company.id, company_name="Acme Corp", full_name="Jane Doe", normalized_name="jane doe",
        job_title="VP Engineering", business_email="jane@acme.com", business_phone="+91-80-1234",
        contact_type=ContactType.BUSINESS_EMAIL, email_status=EmailStatus.VERIFIED_SOURCE,
        verification_status=VerificationStatus.PARTIALLY_VERIFIED, contact_source="Apollo",
        source_type="contact_enrichment", match_score=90, role_match_score=88,
        email_verification_status="VALID", phone_type="BUSINESS_DIRECT",
        phone_verification_status="UNKNOWN", contact_trust_score=55, contact_trust_status="LIKELY",
        is_current=True, data_provenance=DataProvenance.REAL))
    session.commit()


def _load(session):
    res = build_full_intelligence_workbook(session)
    buf = res[0] if isinstance(res, tuple) else res
    return openpyxl.load_workbook(BytesIO(buf.getvalue()))


def test_new_enrichment_columns_present_and_populated(seed_session):
    _seed(seed_session)
    ws = _load(seed_session).active
    headers = [c.value for c in ws[1]]
    for col in ("Role Match", "Email Verification", "Phone Type", "Phone Verification"):
        assert col in headers
    row = dict(zip(headers, [c.value for c in ws[2]]))
    assert row["Target POC"] == "Jane Doe — VP Engineering"
    assert row["Role Match"] == 88
    assert row["Email Verification"] == "VALID"
    assert row["Phone Type"] == "BUSINESS_DIRECT"
    assert row["Contact Trust"] == 55


def test_intelligence_columns_present_and_populated(seed_session):
    """Prompt 50 §36: AI Profile Highlights + POC Status columns."""
    _seed(seed_session)
    ws = _load(seed_session).active
    headers = [c.value for c in ws[1]]
    for col in ("AI Profile Highlights", "POC Status", "POC Last Verified"):
        assert col in headers
    row = dict(zip(headers, [c.value for c in ws[2]]))
    assert row["POC Status"] in ("VERIFIED", "LIKELY", "UNVERIFIED", "STALE", "FORMER")
    # Deterministic highlights are real-data-derived (Hiring Trend present when openings exist).
    assert row["AI Profile Highlights"] and "Hiring Trend" in row["AI Profile Highlights"]


def test_prompt55_columns_present_and_populated(seed_session):
    """Prompt 55 §4: Recommended POC Role, Source URL, Supporting Sources, Evidence."""
    _seed(seed_session)
    ws = _load(seed_session).active
    headers = [c.value for c in ws[1]]
    for col in ("Recommended POC Role", "Source URL", "Supporting Sources", "Evidence"):
        assert col in headers
    row = dict(zip(headers, [c.value for c in ws[2]]))
    # Recommended role is opportunity-aware and present even alongside a real person.
    assert row["Recommended POC Role"]
    # Worksheet name is exactly "Full Intelligence".
    assert ws.title == "Full Intelligence"


def test_fixed_16_column_export_unchanged():
    assert len(FIXED_HEADERS) == 16       # extension never touched the fixed export
