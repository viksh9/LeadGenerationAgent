"""Full intelligence export (Prompt 47, §33/§34).

A SEPARATE, richer export from the fixed 16-column business export (export/excel.py),
which is left completely unchanged. One sheet, one row per REAL company-level lead,
with all useful real business fields — company + legal identity + operating/registered
address + opportunity + POC + per-field source/provenance + Data/Contact Trust.

Real data only: every value is one the application actually holds; blanks stay blank.
No secrets, tokens, or debug payloads are ever exported. Formula-injection-safe cells.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from database.models import Company, DataProvenance, DecisionMaker, Lead, utcnow
from export.excel import (
    _company_id_map,
    _contacts_by_company,
    _poc_text,
    _safe,
    _signal_date,
    _signal_text,
    _source_text,
    _tech_text,
    _enum,
    _PRIORITY_RANK,
    eligible_leads_query,
)
from intelligence.opportunity_view import derive_opportunity_view, opportunity_cell

SHEET_NAME = "Full Intelligence"

# (header, width). One row per company-level real lead. §33 fields that exist in the app.
COLUMNS: list[tuple[str, int]] = [
    ("Lead ID", 10), ("Company Name", 28), ("Legal Company Name", 28), ("Company Number", 22),
    ("Industry", 20), ("Website", 30), ("Operating Address", 34), ("Registered Address", 34),
    ("City", 16), ("State", 16), ("Country", 14), ("India Presence", 14), ("India Entity Type", 18),
    ("Company Status", 14), ("Career URL", 28), ("Contact URL", 28), ("Leadership URL", 28),
    ("LinkedIn Company URL", 30), ("Signal", 34), ("Signal Summary", 40), ("Technology", 30),
    ("Opening Count", 12), ("Opportunity", 22), ("Staffing", 12), ("Estimated Team", 14),
    ("Urgency", 12), ("Score", 8), ("Priority", 10), ("Status", 14),
    ("Target POC", 30), ("POC Role", 22), ("POC LinkedIn", 30), ("Work Email", 28),
    ("Business Phone", 18), ("Data Trust", 10), ("Contact Trust", 12),
    ("Website Source", 24), ("Address Source", 26), ("Registered Address Source", 24),
    ("POC Source", 20), ("Source", 26), ("Registry URL", 30), ("OpenCorporates URL", 34),
    ("Data Trust Verified", 20), ("Signal Date", 16),
]
HEADERS = [c[0] for c in COLUMNS]

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(vertical="top", wrap_text=True)


def _india_presence(company) -> str:
    if company is None or company.india_presence is None:
        return "UNKNOWN"
    return "Yes" if company.india_presence else "No"


def _field_source(session_ev: dict, field: str) -> str:
    e = session_ev.get(field)
    return e.source if e else ""


def full_intelligence_counts(session: Session) -> dict:
    from sqlalchemy import func
    n = int(session.execute(select(func.count(Lead.id)).where(
        Lead.data_provenance == DataProvenance.REAL)).scalar() or 0)
    return {"Leads": n}


def build_full_intelligence_workbook(session: Session, *, now: datetime | None = None
                                     ) -> tuple[BytesIO, dict]:
    now = now or utcnow()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    leads = list(session.execute(eligible_leads_query()).scalars().all())
    leads.sort(key=lambda l: (_PRIORITY_RANK.get(l.lead_priority, -1), float(l.lead_score or 0),
                              (_signal_date(l) or datetime.min)), reverse=True)

    name_map = _company_id_map(session)
    lead_company_id = {l.id: (l.company_id or name_map.get(l.normalized_company_name)) for l in leads}
    companies = {c.id: c for c in session.execute(
        select(Company).where(Company.id.in_({cid for cid in lead_company_id.values() if cid}))
    ).scalars().all()} if lead_company_id else {}
    contacts = _contacts_by_company(session, set(lead_company_id.values()))
    # Field-source lookup (canonical highest-priority) per company.
    from database.models import CompanyFieldEvidence
    ev_by_company: dict[int, dict] = {}
    for cid in {c for c in lead_company_id.values() if c}:
        rows = session.scalars(select(CompanyFieldEvidence).where(
            CompanyFieldEvidence.company_id == cid).order_by(CompanyFieldEvidence.source_priority)).all()
        best: dict = {}
        for e in rows:
            best.setdefault(e.field, e)
        ev_by_company[cid] = best

    row_count = 0
    for lead in leads:
        cid = lead_company_id.get(lead.id)
        company = companies.get(cid) if cid else None
        bundle = contacts.get(cid)
        people = bundle.people if bundle else []
        email_dm = bundle.email_contact if bundle else None
        phone_dm = bundle.phone_contact if bundle else None
        top = people[0] if people else None
        ev = ev_by_company.get(cid, {})
        sig_date = _signal_date(lead)
        oview = derive_opportunity_view(lead, signal_date=sig_date, now=now)
        openings = lead.it_job_count if (lead.it_job_count and lead.it_job_count > 0) else None
        # operating vs registered address kept distinct (§9).
        operating_addr = (company.full_address if company else None) or lead.location_all or lead.location
        registered_addr = company.registered_address if company else None
        shown_sources = [d.contact_source for d in (people[:2] + [d for d in (email_dm, phone_dm) if d])
                         if getattr(d, "contact_source", None)]

        row = [
            lead.id, _safe(lead.company_name),
            _safe(company.legal_name if company else None),
            _safe(company.company_number if company else None),
            _safe((company.industry if company else None) or lead.industry),
            _safe(company.website if company else None),
            _safe(operating_addr), _safe(registered_addr),
            _safe(company.headquarters_city if company else None),
            _safe(company.headquarters_state if company else None),
            _safe(company.headquarters_country if company else None),
            _india_presence(company),
            _safe(company.india_entity_type if company else None),
            _safe(company.company_status if company else None),
            _safe(company.careers_url if company else None),
            _safe(company.contact_url if company else None),
            _safe(company.leadership_url if company else None),
            _safe(company.linkedin_url if company else None),
            _safe(_signal_text(lead)), _safe(lead.signal_description or lead.opportunity_summary),
            _safe(_tech_text(lead)), openings,
            _safe(oview.label), _safe(oview.staffing_need),
            (f"{oview.estimated_team} eng" if oview.estimated_team else None), _safe(oview.urgency),
            round(float(lead.lead_score or 0)), _enum(lead.lead_priority), _enum(lead.status),
            _safe(_poc_text(lead, people)),
            _safe(top.job_title if top else None),
            _safe(top.professional_network_url if top else None),
            _safe(email_dm.business_email if email_dm else None),
            _safe(phone_dm.business_phone if phone_dm else None),
            (company.data_trust_score if company else 0),
            (top.contact_trust_score if top else None),
            _safe(_field_source(ev, "website_url")),
            _safe(_field_source(ev, "address")),
            _safe(_field_source(ev, "registered_address")),
            _safe(top.contact_source if top else None),
            _safe(_source_text(lead, shown_sources)),
            _safe(company.registry_url if company else None),
            _safe(company.opencorporates_url if company else None),
            (company.official_verified_at if company else None),
            sig_date,
        ]
        ws.append(row)
        row_count += 1
        r = row_count + 1
        for col in (2, 3, 7, 8, 19, 20, 21, 30, 41, 42):
            ws.cell(row=r, column=col).alignment = _WRAP
        for date_col in (44, 45):   # Data Trust Verified, Signal Date
            if ws.cell(row=r, column=date_col).value is not None:
                ws.cell(row=r, column=date_col).number_format = "dd-mmm-yyyy"

    if row_count == 0:
        ws.append(["No real lead data available."] + [None] * (len(HEADERS) - 1))

    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(2, row_count + 1)}"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, {"sheet": SHEET_NAME, "total_rows": row_count, "columns": len(HEADERS),
                    "generated_at": now.isoformat()}
